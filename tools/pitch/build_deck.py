# -*- coding: utf-8 -*-
"""路演 PPTX 生成：反向肉鸽《锅巴必须死》—— 锅巴队。

设计取向：深色底 + 血红/紫罗兰两色，一页一件事，台上照着念就能讲完。
所有文案都来自 docs/GAME_DESIGN.md 与 README，不临场编卖点。
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

W, H = Inches(13.333), Inches(7.5)  # 16:9

BG = RGBColor(0x0B, 0x09, 0x12)
INK = RGBColor(0xF2, 0xEC, 0xFF)
DIM = RGBColor(0x9A, 0x92, 0xB5)
BLOOD = RGBColor(0xE0, 0x2B, 0x45)
VIOLET = RGBColor(0x9B, 0x6B, 0xFF)
GOLD = RGBColor(0xFF, 0xB3, 0x47)
CARD = RGBColor(0x16, 0x13, 0x22)

FONT = "微软雅黑"


def slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    return s


def text(s, x, y, w, h, body, size, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.25):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = body.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = spacing
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.name = FONT
        p.font.color.rgb = color
    return box


def rule(s, x, y, w, color=BLOOD, thickness=Inches(0.035)):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, thickness)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def card(s, x, y, w, h, title, lines, accent=VIOLET):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.color.rgb = RGBColor(0x2E, 0x27, 0x44)
    box.line.width = Pt(1)
    rule(s, x + Inches(0.3), y + Inches(0.28), Inches(0.5), accent, Inches(0.045))
    text(s, x + Inches(0.3), y + Inches(0.38), w - Inches(0.6), Inches(0.5), title, 19, INK, True)
    text(s, x + Inches(0.3), y + Inches(0.95), w - Inches(0.6), h - Inches(1.1), lines, 13.5, DIM)


def bullet(s, x, y, w, rows, size=16, gap=Inches(0.62)):
    for i, (head, body) in enumerate(rows):
        yy = y + gap * i
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x, yy + Inches(0.09), Inches(0.1), Inches(0.1))
        dot.fill.solid()
        dot.fill.fore_color.rgb = BLOOD
        dot.line.fill.background()
        text(s, x + Inches(0.25), yy, w, Inches(0.4), head, size, INK, True)
        if body:
            text(s, x + Inches(0.25), yy + Inches(0.34), w, Inches(0.4), body, 12.5, DIM)


prs = Presentation()
prs.slide_width, prs.slide_height = W, H

# ---------------------------------------------------------------- 1 封面
s = slide(prs)
rule(s, Inches(1.0), Inches(1.55), Inches(0.9), BLOOD)
text(s, Inches(1.0), Inches(1.75), Inches(10), Inches(0.5), "反向肉鸽 · 2D 俯视角 · 单局 5 分钟 · 浏览器直接玩", 15, DIM)
text(s, Inches(0.95), Inches(2.3), Inches(11.6), Inches(1.6), "锅巴必须死", 78, INK, True)
text(s, Inches(1.0), Inches(3.85), Inches(11.2), Inches(0.9),
     "你是魔王。你连身体都没有。\n把魔物摆好，在五分钟内杀掉来讨伐你的勇者「锅巴」。", 20, VIOLET)
rule(s, Inches(1.0), Inches(5.25), Inches(3.2), RGBColor(0x33, 0x2C, 0x48), Inches(0.02))
text(s, Inches(1.0), Inches(5.45), Inches(6), Inches(0.4), "锅巴队 · 郭峰 / 周玉通 / 周家名", 16, INK, True)
text(s, Inches(1.0), Inches(5.85), Inches(6), Inches(0.4), "WaytoAGI 1024 切磋大会 · 游戏赛道", 13, DIM)

# ---------------------------------------------------------------- 2 一句话（终选简介，逐字）
s = slide(prs)
text(s, Inches(1.0), Inches(0.75), Inches(11), Inches(0.5), "一句话", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.2), Inches(11.2), Inches(1.7),
     "普通肉鸽是「你去杀怪」。\n这一款反过来——怪是你的，勇者是来杀你的。", 32, INK, True)
rule(s, Inches(1.0), Inches(3.05), Inches(4.2), VIOLET, Inches(0.03))
# 这一页就是提交表单里那段简介，逐字一致，台上直接念
text(s, Inches(1.0), Inches(3.35), Inches(11.3), Inches(3.0),
     "你是魔王。你连身体都没有，所以画面上看不到你：镜头永远跟着来讨伐你的勇者「锅巴」。\n"
     "你能做的只有一件事——在他必经之路上，提前把魔物摆好。\n"
     "他杀你的怪会变强，你每 30 秒拿一条「七宗罪」权能也会让你的怪变强。\n"
     "五分钟内把他留下，或者看着他推平你的王座。", 19, INK, spacing=1.55)

# ---------------------------------------------------------------- 3 玩法
s = slide(prs)
text(s, Inches(1.0), Inches(0.6), Inches(11), Inches(0.5), "怎么玩", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.0), Inches(11), Inches(0.7), "30 秒一个节拍，一局 10 拍", 30, INK, True)
card(s, Inches(1.0), Inches(2.0), Inches(3.5), Inches(2.5), "① 摆怪",
     "5 张魔物卡，星级 1–10。\n星数上限随时间上涨，\n所以牌面会越铺越大。", BLOOD)
card(s, Inches(4.8), Inches(2.0), Inches(3.5), Inches(2.5), "② 收权能",
     "每 30 秒三选一「七宗罪」：\n暴食、色欲、贪婪……\n每条都在改这局的规则。", VIOLET)
card(s, Inches(8.6), Inches(2.0), Inches(3.5), Inches(2.5), "③ 看他成长",
     "他杀怪升级、学七美德、\n换技能。你是看着他变强，\n再想办法把他按回去。", GOLD)
text(s, Inches(1.0), Inches(4.85), Inches(11.2), Inches(1.4),
     "五张卡＝五种走位：史莱姆堵路 / 蝙蝠群穿行 / 哥布林远程消耗 / 石像鬼卡弹道 / 巨大蜘蛛吐丝近咬还生小蜘蛛。\n"
     "没有经济系统、没有抽卡——只有「提前布置」这一件事。", 15, DIM)

# ---------------------------------------------------------------- 4 记忆点
s = slide(prs)
text(s, Inches(1.0), Inches(0.7), Inches(11), Inches(0.5), "记忆点", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.1), Inches(11), Inches(0.7), "三个别人抄不走的地方", 30, INK, True)
bullet(s, Inches(1.0), Inches(2.2), Inches(11),
       [("① 看不见的魔王", "玩家就是魔王，但画面上没有魔王——只有他控制的怪，和一路压过来的锅巴。"),
        ("② 相机跟着敌人", "镜头永远跟着锅巴，玩家是在「被讨伐方的视角」里做决策。"),
        ("③ 双向成长的对赌", "他每杀一只怪就变强一点，你每 30 秒拿一条权能；两边曲线谁先拐起来，这局就归谁。")],
       size=19)

# ---------------------------------------------------------------- 5 锅巴
s = slide(prs)
text(s, Inches(1.0), Inches(0.7), Inches(11), Inches(0.5), "对手", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.1), Inches(11), Inches(0.7), "锅巴不是一个靶子，是一个会读局的对手", 28, INK, True)
bullet(s, Inches(1.0), Inches(2.15), Inches(11),
       [("7 个主动技能", "点射 / 散射 / 剑气 / 落雷 / 重击 / 干粮 / 翻滚，放技能时站定，被贴脸会翻走。"),
        ("7 条美德天赋", "你选罪，他长德：谦逊、贞洁、节制……重复获得会升级，越打越像精英。"),
        ("感知与状态机", "10Hz 感知 + 目标迟滞 + 追击/撤退/蓄力/暴怒四态，不是直线冲脸的木头人。")],
       size=18)

# ---------------------------------------------------------------- 6 美术与音频
s = slide(prs)
text(s, Inches(1.0), Inches(0.7), Inches(11), Inches(0.5), "怎么做出来的", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.1), Inches(11), Inches(0.7), "AI 出图 + 零音频文件", 30, INK, True)
card(s, Inches(1.0), Inches(2.15), Inches(5.4), Inches(2.9), "美术：图像模型生成",
     "怪物、勇者、场景、UI 图标全部由图像模型出图，\n统一暗紫血红调，尺寸与锚点对齐后进 Atlas。\n人工只做筛选与修边。", VIOLET)
card(s, Inches(6.7), Inches(2.15), Inches(5.4), Inches(2.9), "音乐音效：100% 程序合成",
     "没有一个音频文件——全部用 WebAudio 现场合成：\n鼓点、法术、命中、升级、败北都是振荡器与噪声。\n加载体积为零。", GOLD)
text(s, Inches(1.0), Inches(5.35), Inches(11.2), Inches(1.2),
     "工程：TypeScript(strict) + PixiJS 8 + 固定步长 1/60 逻辑。\n"
     "交付物是一个静态站点，点开链接就能玩，不用装、不用登录。", 15, DIM)

# ---------------------------------------------------------------- 7 宣传图
s = slide(prs)
pic = s.shapes.add_picture("art/pitch/poster-4x3.png", 0, 0, height=H)
pic.left = int((W - pic.width) / 2)

# ---------------------------------------------------------------- 8 现在能玩
s = slide(prs)
text(s, Inches(1.0), Inches(0.7), Inches(11), Inches(0.5), "现状", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.1), Inches(11), Inches(0.7), "不是 Demo：完整一局已经跑通", 30, INK, True)
bullet(s, Inches(1.0), Inches(2.15), Inches(11),
       [("完整 5 分钟一局", "胜负判定、结算数据、再来一局全部闭环；新手教程进局先讲一遍。"),
        ("在线可玩", "https://gove2004.github.io/20261024/ —— 手机横竖屏都能开。"),
        ("质量门", "冒烟测试 70/70 通过，线上探针 LIVE PASS；20 条规则单测全绿。")],
       size=18)

# ---------------------------------------------------------------- 9 团队
s = slide(prs)
text(s, Inches(1.0), Inches(1.4), Inches(11), Inches(0.5), "锅巴队", 14, BLOOD, True)
text(s, Inches(1.0), Inches(1.85), Inches(11), Inches(0.8), "郭峰 · 周玉通 · 周家名", 40, INK, True)
rule(s, Inches(1.0), Inches(2.95), Inches(3.0), VIOLET, Inches(0.03))
text(s, Inches(1.0), Inches(3.25), Inches(11.2), Inches(2.2),
     "三个人，五天，从策划案写到上线：玩法与数值、战斗 AI、渲染与 UI、美术管线、自动化验证各有人盯。\n"
     "《锅巴必须死》—— 2026-10-25 · WaytoAGI 1024 切磋大会。", 17, DIM)
text(s, Inches(1.0), Inches(5.9), Inches(6), Inches(0.5), "谢谢，来玩一局吧：你现在是魔王了。", 20, VIOLET, True)

out = "art/pitch/锅巴必须死-路演.pptx"
prs.save(out)
print("saved", out)
